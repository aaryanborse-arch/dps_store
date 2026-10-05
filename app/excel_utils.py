from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter
from app.pdf_utils import amount_to_words, STORE_NAME, STORE_ADDRESS, STORE_GSTIN

thin = Side(style='thin', color='999999')
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
HEADER_FILL = PatternFill(start_color='343A40', end_color='343A40', fill_type='solid')
HEADER_FONT = Font(color='FFFFFF', bold=True)
BOLD = Font(bold=True)


def generate_bill_excel(bill):
    wb = Workbook()
    ws = wb.active
    ws.title = "Tax Invoice"

    col_count = 9
    for i in range(1, col_count + 1):
        ws.column_dimensions[get_column_letter(i)].width = 14

    row = 1
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=col_count)
    ws.cell(row=row, column=1, value=STORE_NAME).font = Font(bold=True, size=14)
    ws.cell(row=row, column=1).alignment = Alignment(horizontal='center')
    row += 1

    extra_line = bill.city.extra_counter_line if bill.city and bill.city.extra_counter_line else ''
    for line in [STORE_ADDRESS, extra_line, f"GSTIN NO. {STORE_GSTIN}"]:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=col_count)
        ws.cell(row=row, column=1, value=line).alignment = Alignment(horizontal='center')
        row += 1

    row += 1
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=col_count)
    ws.cell(row=row, column=1, value="TAX INVOICE").font = Font(bold=True, size=13, underline='single')
    ws.cell(row=row, column=1).alignment = Alignment(horizontal='center')
    row += 2

    issued_by = bill.created_by.username if bill.created_by else '-'
    meta_rows = [
        ('Invoice No.:', bill.bill_number, 'Date:', bill.created_at.strftime('%d-%m-%Y')),
        ('Student Name:', bill.student_name or '-', 'Class:', bill.student_class or '-'),
        ('Buyer:', bill.buyer_name, 'Payment:', bill.payment_method or '-'),
        ('Issued By:', issued_by, '', ''),
    ]
    for a, b, c, d in meta_rows:
        ws.cell(row=row, column=1, value=a).font = BOLD
        ws.cell(row=row, column=2, value=b)
        ws.cell(row=row, column=4, value=c).font = BOLD
        ws.cell(row=row, column=5, value=d)
        row += 1

    row += 1
    headers = ['S.No', 'Particulars', 'Publisher', 'Subject', 'HSN', 'Qty', 'Rate', 'GST%', 'Amount']
    for col, h in enumerate(headers, start=1):
        c = ws.cell(row=row, column=col, value=h)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.border = BORDER
        c.alignment = Alignment(horizontal='center')
    row += 1

    hsn_summary = {}
    total_qty = 0
    grand_taxable = 0.0
    grand_tax = 0.0

    for idx, item in enumerate(bill.items, start=1):
        gst_display = 'Exempt' if item.is_gst_exempt else f'{item.gst_percent}%'
        base_amount = item.subtotal
        gst_amount = 0.0 if item.is_gst_exempt else round(base_amount * item.gst_percent / 100, 2)
        line_total = base_amount + gst_amount

        values = [idx, item.product_name, item.publisher_name or '-', item.subject or '-',
                  item.hsn_number, item.quantity, item.price, gst_display, round(line_total, 2)]
        for col, v in enumerate(values, start=1):
            c = ws.cell(row=row, column=col, value=v)
            c.border = BORDER
        row += 1

        total_qty += item.quantity
        grand_taxable += base_amount
        grand_tax += gst_amount

        key = (item.hsn_number, item.gst_percent, item.is_gst_exempt)
        if key not in hsn_summary:
            hsn_summary[key] = {'taxable': 0.0, 'gst': 0.0}
        hsn_summary[key]['taxable'] += base_amount
        hsn_summary[key]['gst'] += gst_amount

    grand_total = grand_taxable + grand_tax

    total_row_values = ['', '', '', '', 'TOTAL', total_qty, '', '', round(grand_total, 2)]
    for col, v in enumerate(total_row_values, start=1):
        c = ws.cell(row=row, column=col, value=v)
        c.font = BOLD
        c.border = BORDER
    row += 2

    hsn_headers = ['HSN/SAC', 'Taxable Value', 'Central Tax Rate', 'Central Tax Amt', 'State Tax Rate', 'State Tax Amt', 'Total Tax']
    for col, h in enumerate(hsn_headers, start=1):
        c = ws.cell(row=row, column=col, value=h)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.border = BORDER
        c.alignment = Alignment(horizontal='center')
    row += 1

    for (hsn, gst_percent, is_exempt), vals in hsn_summary.items():
        half_rate = 0 if is_exempt else gst_percent / 2
        half_amt = vals['gst'] / 2
        values = [hsn, round(vals['taxable'], 2),
                  'Exempt' if is_exempt else f'{half_rate}%', round(half_amt, 2),
                  'Exempt' if is_exempt else f'{half_rate}%', round(half_amt, 2),
                  round(vals['gst'], 2)]
        for col, v in enumerate(values, start=1):
            c = ws.cell(row=row, column=col, value=v)
            c.border = BORDER
        row += 1

    hsn_total_values = ['TOTAL', round(grand_taxable, 2), '', round(grand_tax / 2, 2), '', round(grand_tax / 2, 2), round(grand_tax, 2)]
    for col, v in enumerate(hsn_total_values, start=1):
        c = ws.cell(row=row, column=col, value=v)
        c.font = BOLD
        c.border = BORDER
    row += 2

    ws.cell(row=row, column=1, value='Amount Chargeable [In Words]').font = BOLD
    ws.cell(row=row, column=3, value='GRAND TOTAL').font = Font(bold=True, size=12)
    ws.cell(row=row, column=6, value=round(grand_total, 2)).font = Font(bold=True, size=12)
    row += 1
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
    ws.cell(row=row, column=1, value=amount_to_words(grand_total))
    ws.cell(row=row, column=9, value='E. & O. E.').font = BOLD
    row += 3

    ws.cell(row=row, column=1, value='Books Checked By')
    ws.cell(row=row, column=6, value=f'For, {STORE_NAME}').font = BOLD

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer

def generate_day_wise_excel(data, city_label):
    wb = Workbook()
    ws = wb.active
    ws.title = "Day Wise Sale Report"

    rates = data['rates']
    header = ['Sr.', 'Date', 'Invoice No.', 'Student Name', 'Class', 'Payment Mode', 'Exempt Amt']
    for r in rates:
        header.append(f'Taxable @{r}%')
    for r in rates:
        header.append(f'GST @{r}%')
    header.append('Grand Total')

    col_count = len(header)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=col_count)
    date_label = data['start'].strftime('%d-%m-%Y') if data['start'].date() == data['end'].date() else f"{data['start'].strftime('%d-%m-%Y')} to {data['end'].strftime('%d-%m-%Y')}"
    title_cell = ws.cell(row=1, column=1, value=f"{STORE_NAME} - SALES REPORT - {date_label} ({city_label})")
    title_cell.font = Font(bold=True, size=13)
    title_cell.alignment = Alignment(horizontal='center')

    header_row = 3
    for col, h in enumerate(header, start=1):
        c = ws.cell(row=header_row, column=col, value=h)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.border = BORDER
        c.alignment = Alignment(horizontal='center', wrap_text=True)

    data_fill = PatternFill(start_color='E7F3EA', end_color='E7F3EA', fill_type='solid')
    total_col_fill = PatternFill(start_color='FBE0CE', end_color='FBE0CE', fill_type='solid')
    total_row_fill = PatternFill(start_color='D9A441', end_color='D9A441', fill_type='solid')

    row = header_row + 1
    for r in data['rows']:
        values = [r['sr_no'], r['date'].strftime('%d-%m-%Y'), r['bill_number'], r['student_name'],
                  r['student_class'], r['payment_mode'], r['exempt_amt']]
        for rate in rates:
            values.append(r['taxable_by_rate'].get(rate, 0))
        for rate in rates:
            values.append(r['gst_by_rate'].get(rate, 0))
        values.append(r['grand_total'])

        for col, v in enumerate(values, start=1):
            c = ws.cell(row=row, column=col, value=v)
            c.border = BORDER
            c.fill = total_col_fill if col == col_count else data_fill
        row += 1

    total_values = ['', '', '', '', '', 'TOTAL', data['total_exempt']]
    for rate in rates:
        total_values.append(data['total_taxable_by_rate'].get(rate, 0))
    for rate in rates:
        total_values.append(data['total_gst_by_rate'].get(rate, 0))
    total_values.append(data['total_grand'])

    for col, v in enumerate(total_values, start=1):
        c = ws.cell(row=row, column=col, value=v)
        c.font = BOLD
        c.border = BORDER
        c.fill = total_row_fill

    for i in range(1, col_count + 1):
        ws.column_dimensions[get_column_letter(i)].width = 16

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer

def generate_report_excel(period, ref_date, data, city_label):
    wb = Workbook()
    ws = wb.active
    ws.title = "Stock Report"

    ws.merge_cells('A1:F1')
    title_cell = ws.cell(row=1, column=1, value=f"{STORE_NAME} - {period.capitalize()} Stock Report - {city_label}")
    title_cell.font = Font(bold=True, size=13)
    title_cell.alignment = Alignment(horizontal='center')

    ws.merge_cells('A2:F2')
    ws.cell(row=2, column=1, value=f"Period: {data['start'].strftime('%d-%m-%Y')} to {data['end'].strftime('%d-%m-%Y')}").alignment = Alignment(horizontal='center')

    row = 4
    ws.cell(row=row, column=1, value='Total Bills:').font = BOLD
    ws.cell(row=row, column=2, value=data['total_bills'])
    ws.cell(row=row, column=4, value='Total Amount:').font = BOLD
    ws.cell(row=row, column=5, value=round(data['total_amount'], 2))
    row += 2

    header_row = row
    headers = ['Sr.No', 'Product', 'Year', 'MRP', 'Opening Stock', 'Inward Stock', 'Current Stock', 'Sales', 'Closing Stock']
    for col, h in enumerate(headers, start=1):
        c = ws.cell(row=header_row, column=col, value=h)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.border = BORDER
        c.alignment = Alignment(horizontal='center')
    row += 1

    for idx, s in enumerate(data['stock_overview'], start=1):
        values = [idx, s['name'], s.get('year') or '-', s.get('price', 0), s['opening_stock'], s['inward_stock'], s['current_stock'], s['sales'], s['closing_stock']]
        for col, v in enumerate(values, start=1):
            c = ws.cell(row=row, column=col, value=v)
            c.border = BORDER
        row += 1

    widths = [30, 14, 14, 10, 14, 14]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer

def generate_gst_report_excel(period, start_date, data, city_label):
    wb = Workbook()
    ws = wb.active
    ws.title = "GST Report"

    ws.merge_cells('A1:G1')
    title_cell = ws.cell(row=1, column=1, value=f"{STORE_NAME} - GST REPORT - {city_label} ({period.capitalize()})")
    title_cell.font = Font(bold=True, size=13)
    title_cell.alignment = Alignment(horizontal='center')

    ws.merge_cells('A2:G2')
    ws.cell(row=2, column=1, value=f"Period: {data['start'].strftime('%d-%m-%Y')} to {data['end'].strftime('%d-%m-%Y')}").alignment = Alignment(horizontal='center')

    row = 4
    summary = [
        ('Taxable Value', data['total_taxable']), ('Total CGST', data['total_cgst']),
        ('Total SGST', data['total_sgst']), ('Grand Total', data['grand_total'])
    ]
    for label, val in summary:
        ws.cell(row=row, column=1, value=label).font = BOLD
        ws.cell(row=row, column=2, value=round(val, 2))
        row += 1
    row += 1

    header_row = row
    headers = ['Bill No.', 'Date', 'Taxable Value', 'CGST', 'SGST', 'Total GST', 'Grand Total']
    for col, h in enumerate(headers, start=1):
        c = ws.cell(row=header_row, column=col, value=h)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.border = BORDER
        c.alignment = Alignment(horizontal='center')
    row += 1

    for r in data['bill_rows']:
        values = [r['bill_number'], r['date'].strftime('%d-%m-%Y %H:%M'), round(r['taxable'], 2),
                  round(r['cgst'], 2), round(r['sgst'], 2), round(r['total_gst'], 2), round(r['grand_total'], 2)]
        for col, v in enumerate(values, start=1):
            c = ws.cell(row=row, column=col, value=v)
            c.border = BORDER
        row += 1

    widths = [16, 18, 14, 12, 12, 12, 14]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer