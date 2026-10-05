from io import BytesIO
from flask import Blueprint, render_template, redirect, url_for, request, flash, send_file
from flask_login import login_required, current_user
from functools import wraps
from openpyxl import Workbook, load_workbook
from app import db
from app.models import StudentRecord, City
from app.class_options import CLASS_OPTIONS
from app.class_options import CLASS_OPTIONS, normalize_class_value
from app.inventory import store_staff_required
from app.city_utils import get_active_city_id

students_bp = Blueprint('students_bp', __name__, url_prefix='/students')


def super_admin_required(f):
    @wraps(f)
    def decorated(*args, **kwargs):
        if not current_user.is_authenticated or current_user.role != 'super_admin':
            flash('Access denied — Super Admin only')
            return redirect(url_for('auth.dashboard'))
        return f(*args, **kwargs)
    return decorated


@students_bp.route('/template')
@login_required
@super_admin_required
def download_template():
    wb = Workbook()
    ws = wb.active
    ws.append(['Admission Number', 'Student Name', 'Class'])
    ws.append(['456', 'Advait Javkar', '1'])
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return send_file(buffer, as_attachment=True, download_name='student_list_template.xlsx',
                      mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


@students_bp.route('/', methods=['GET', 'POST'])
@login_required
@super_admin_required
def manage_students():
    all_cities = City.query.order_by(City.name).all()

    city_id_param = request.args.get('city_id') or request.form.get('city_id')
    city_id = int(city_id_param) if city_id_param else (all_cities[0].id if all_cities else None)

    if request.method == 'POST':
        if not city_id:
            flash('Select a school first')
            return redirect(url_for('students_bp.manage_students'))

        file = request.files.get('excel_file')
        if not file or file.filename == '':
            flash('Please choose an Excel file')
            return redirect(url_for('students_bp.manage_students', city_id=city_id))

        try:
            wb = load_workbook(file, data_only=True)
        except Exception:
            flash('Could not read that file')
            return redirect(url_for('students_bp.manage_students', city_id=city_id))

        if 'Students' in wb.sheetnames:
            ws = wb['Students']
        else:
            ws = wb.active
            flash('Note: no sheet named "Students" found — used the first sheet instead. Double-check the results below.')

        def find_col(header_row, keywords):
            for idx, val in enumerate(header_row):
                if any(kw in val for kw in keywords):
                    return idx
            return None

        header_row_idx = None
        adm_col = name_col = class_col = None

        for row_idx in range(1, 11):
            candidate = [str(c.value).strip().lower() if c.value else '' for c in ws[row_idx]]
            a = find_col(candidate, ['admission no', 'admission number', 'adm no', 'adm. no'])
            n = find_col(candidate, ['student name', 'name'])
            c = find_col(candidate, ['class'])
            if a is not None and n is not None and c is not None:
                header_row_idx = row_idx
                adm_col, name_col, class_col = a, n, c
                break

        if header_row_idx is None:
            flash('Could not find a row with Admission No, Student Name and Class column headers in the first 10 rows. Make sure those exact labels appear together somewhere near the top of the file.')
            return redirect(url_for('students_bp.manage_students', city_id=city_id))

        def clean_admission_no(val):
            text = str(val).strip()
            if text.endswith('.0'):
                text = text[:-2]
            return text

        full_replace = request.form.get('full_replace') == 'on'

        # load all existing students for this school ONCE, instead of querying per row
        existing_students = {s.admission_number: s for s in StudentRecord.query.filter_by(city_id=city_id).all()}

        imported = 0
        seen_admission_numbers = set()
        new_records = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or adm_col >= len(row) or not row[adm_col]:
                continue
            adm_no = clean_admission_no(row[adm_col])
            name = str(row[name_col]).strip() if name_col < len(row) and row[name_col] else None
            cls = str(row[class_col]).strip() if class_col < len(row) and row[class_col] else None
            if not adm_no or not name or not cls:
                continue

            seen_admission_numbers.add(adm_no)

            existing = existing_students.get(adm_no)
            if existing:
                existing.student_name = name
                existing.student_class = cls
            else:
                new_records.append(StudentRecord(city_id=city_id, admission_number=adm_no, student_name=name, student_class=cls))
            imported += 1

        if new_records:
            db.session.bulk_save_objects(new_records)

        removed = 0
        if full_replace and seen_admission_numbers:
            students_to_remove = StudentRecord.query.filter(
                StudentRecord.city_id == city_id,
                ~StudentRecord.admission_number.in_(seen_admission_numbers)
            ).all()
            removed = len(students_to_remove)
            for s in students_to_remove:
                db.session.delete(s)

        db.session.commit()

        if full_replace:
            flash(f'Full list replaced: {imported} student(s) added/updated, {removed} old record(s) removed (not present in this file)')
        else:
            flash(f'Imported/updated {imported} student record(s)')
        return redirect(url_for('students_bp.manage_students', city_id=city_id))

    records = StudentRecord.query.filter_by(city_id=city_id).order_by(StudentRecord.student_name).all() if city_id else []
    selected_city = City.query.get(city_id) if city_id else None

    return render_template('students/manage.html', records=records, class_options=CLASS_OPTIONS,
                            all_cities=all_cities, selected_city=selected_city)


@students_bp.route('/add', methods=['POST'])
@login_required
@super_admin_required
def add_student():
    city_id = request.form.get('city_id')
    admission_number = request.form.get('admission_number', '').strip()
    student_name = request.form.get('student_name', '').strip()
    student_class = request.form.get('student_class', '').strip()

    if not city_id or not admission_number or not student_name or not student_class:
        flash('School, admission number, student name and class are all required')
        return redirect(url_for('students_bp.manage_students'))

    city_id = int(city_id)

    existing = StudentRecord.query.filter_by(city_id=city_id, admission_number=admission_number).first()
    if existing:
        existing.student_name = student_name
        existing.student_class = student_class
        flash(f'Updated existing record for admission no. {admission_number}')
    else:
        db.session.add(StudentRecord(city_id=city_id, admission_number=admission_number,
                                      student_name=student_name, student_class=student_class))
        flash(f'Added student "{student_name}" (Admission No. {admission_number})')

    db.session.commit()
    return redirect(url_for('students_bp.manage_students', city_id=city_id))


@students_bp.route('/<int:student_id>/delete', methods=['POST'])
@login_required
@super_admin_required
def delete_student(student_id):
    student = StudentRecord.query.get_or_404(student_id)
    city_id = student.city_id
    name = student.student_name
    db.session.delete(student)
    db.session.commit()
    flash(f'Removed student "{name}"')
    return redirect(url_for('students_bp.manage_students', city_id=city_id))


@students_bp.route('/delete-all', methods=['POST'])
@login_required
@super_admin_required
def delete_all_students():
    city_id = request.form.get('city_id')
    if not city_id:
        flash('No school specified')
        return redirect(url_for('students_bp.manage_students'))

    city_id = int(city_id)
    count = StudentRecord.query.filter_by(city_id=city_id).delete()
    db.session.commit()
    flash(f'Removed all {count} student record(s) for this school')
    return redirect(url_for('students_bp.manage_students', city_id=city_id))

@students_bp.route('/staff-add', methods=['GET', 'POST'])
@login_required
@store_staff_required
def staff_add_student():
    city_id = get_active_city_id()

    if request.method == 'POST':
        admission_number = request.form.get('admission_number', '').strip()
        student_name = request.form.get('student_name', '').strip()
        student_class = request.form.get('student_class', '').strip()

        if not admission_number or not student_name or not student_class:
            flash('Admission number, student name and class are all required')
            return redirect(url_for('students_bp.staff_add_student'))

        existing = StudentRecord.query.filter_by(city_id=city_id, admission_number=admission_number).first()
        if existing:
            existing.student_name = student_name
            existing.student_class = student_class
            flash(f'Updated existing record for admission no. {admission_number}')
        else:
            db.session.add(StudentRecord(city_id=city_id, admission_number=admission_number,
                                          student_name=student_name, student_class=student_class))
            flash(f'Added student "{student_name}" (Admission No. {admission_number})')

        db.session.commit()
        return redirect(url_for('students_bp.staff_add_student'))

    records = StudentRecord.query.filter_by(city_id=city_id).order_by(StudentRecord.student_name).all()
    return render_template('students/staff_add.html', records=records, class_options=CLASS_OPTIONS)