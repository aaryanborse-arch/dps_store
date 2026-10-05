import re
from openpyxl import load_workbook
from flask import Blueprint, render_template, redirect, url_for, request, flash
from flask_login import login_required, current_user
from app import db
from app.models import Product, StockLog, City
from app.inventory import store_staff_required, parse_gst
from app.city_utils import get_active_city_id
from app.class_options import CLASS_OPTIONS, get_class_label

catalog_import = Blueprint('catalog_import', __name__, url_prefix='/inventory/catalog-import')

# maps common sheet-name spellings to our internal class keys
SHEET_NAME_MAP = {
    'nursery': 'Nursery', 'lkg': 'Pre-Prep', 'ukg': 'Prep',
    'i': '1', 'ii': '2', 'iii': '3', 'iv': '4', 'v': '5',
    'vi': '6', 'vii': '7', 'viii': '8', 'ix': '9', 'x': '10',
    'xi-math': '11-math', 'xi-com': '11-commerce', 'xi-commerce': '11-commerce',
    'xi-bio': '11-bio', 'xi-hum': '11-humanities', 'xi-humanities': '11-humanities',
    'xii-math': '12-math', 'xii-com.': '12-commerce', 'xii-com': '12-commerce', 'xii-commerce': '12-commerce',
    'xii-bio': '12-bio', 'xii-hum.': '12-humanities', 'xii-hum': '12-humanities', 'xii-humanities': '12-humanities',
}

HEADER_KEYWORDS = {
    'name': ['name of the book', 'books name', 'particulars', 'book name'],
    'subject': ['sub', 'subject'],
    'publisher': ['publisher', 'pub'],
    'qty': ['qty', 'quantity'],
    'gst': ['gst rate', 'gst'],
    'rate': ['rate'],
    'amount': ['amount'],
}


def map_sheet_name_to_class(sheet_name):
    key = re.sub(r'\s+', '', sheet_name.strip().lower())
    return SHEET_NAME_MAP.get(key)


def find_header_row(ws, max_scan=10):
    for row_idx in range(1, max_scan + 1):
        row_values = [str(c.value).strip().lower() if c.value else '' for c in ws[row_idx]]
        joined = ' '.join(row_values)
        if any(k in joined for k in ['qty', 'quantity']) and any(k in joined for k in ['rate', 'amount']):
            return row_idx, row_values
    return None, None


def map_columns(header_values):
    col_map = {}
    for key, keywords in HEADER_KEYWORDS.items():
        for idx, val in enumerate(header_values):
            if any(kw in val for kw in keywords):
                col_map[key] = idx
                break
    return col_map


def parse_sheet(ws, class_key):
    header_row_idx, header_values = find_header_row(ws)
    if header_row_idx is None:
        return []

    col_map = map_columns(header_values)
    if 'name' not in col_map:
        return []

    rows = []
    is_elective_zone = False

    for row in ws.iter_rows(min_row=header_row_idx + 1, values_only=True):
        if row is None:
            continue

        def get_val(key):
            idx = col_map.get(key)
            if idx is None or idx >= len(row):
                return None
            return row[idx]

        name_val = get_val('name')
        subject_val = get_val('subject')

        # detect the "OPTIONAL" marker row — everything after it in this sheet is elective
        row_text = ' '.join(str(v).strip().upper() for v in row if v)
        if 'OPTIONAL' in row_text and ('TOTAL' in row_text or not name_val):
            is_elective_zone = True
            continue

        if not name_val or not str(name_val).strip():
            continue
        name = str(name_val).strip()
        if name.upper() in ('TOTAL', 'GRAND TOTAL'):
            continue

        qty_val = get_val('qty')
        gst_val = get_val('gst')
        rate_val = get_val('rate')
        amount_val = get_val('amount')
        publisher_val = get_val('publisher')

        try:
            qty = int(float(qty_val)) if qty_val not in (None, '') else 1
        except (ValueError, TypeError):
            qty = 1

        try:
            if rate_val not in (None, ''):
                price = float(rate_val)
            elif amount_val not in (None, ''):
                price = float(amount_val)
            else:
                price = 0.0
        except (ValueError, TypeError):
            price = 0.0

        gst_raw = str(gst_val).strip() if gst_val not in (None, '') else 'exempt'
        try:
            gst_percent, is_exempt = parse_gst(gst_raw)
        except (ValueError, TypeError):
            gst_percent, is_exempt = 0.0, True

        rows.append({
            'name': name,
            'subject': str(subject_val).strip() if subject_val else '',
            'publisher': str(publisher_val).strip() if publisher_val else '',
            'price': round(price, 2),
            'gst_percent': gst_percent,
            'is_exempt': is_exempt,
            'quantity': qty,
            'class_key': class_key,
            'is_elective': is_elective_zone,
        })

    return rows


@catalog_import.route('/upload', methods=['GET', 'POST'])
@login_required
@store_staff_required
def upload():
    if request.method == 'POST':
        file = request.files.get('excel_file')
        if not file or file.filename == '':
            flash('Please choose an Excel file')
            return redirect(url_for('catalog_import.upload'))

        try:
            wb = load_workbook(file, data_only=True)
        except Exception:
            flash('Could not read that file')
            return redirect(url_for('catalog_import.upload'))

        all_rows = []
        unmatched_sheets = []

        for sheet_name in wb.sheetnames:
            class_key = map_sheet_name_to_class(sheet_name)
            if class_key is None:
                unmatched_sheets.append(sheet_name)
                continue
            ws = wb[sheet_name]
            rows = parse_sheet(ws, class_key)
            all_rows.extend(rows)

        if not all_rows:
            flash('Could not extract any items — check the file matches the expected per-class-sheet format')
            return redirect(url_for('catalog_import.upload'))

        if unmatched_sheets:
            flash(f'Skipped sheets with unrecognized names (not matched to a class): {", ".join(unmatched_sheets)}')

        return render_template('catalog_import/review.html', rows=all_rows, class_options=CLASS_OPTIONS)

    return render_template('catalog_import/upload.html')


@catalog_import.route('/confirm', methods=['POST'])
@login_required
@store_staff_required
def confirm():
    city_id = get_active_city_id()

    names = request.form.getlist('name')
    subjects = request.form.getlist('subject')
    publishers = request.form.getlist('publisher')
    class_keys = request.form.getlist('class_key')
    gst_percents = request.form.getlist('gst_percent')
    prices = request.form.getlist('price')
    quantities = request.form.getlist('quantity')
    electives = request.form.getlist('is_elective')  # only checked boxes appear, matched by index via hidden mirror

    elective_flags = request.form.getlist('is_elective_flag')  # 'true'/'false' per row, always present

    clear_first = request.form.get('clear_existing') == 'on'

    if clear_first:
        affected_classes = set(k for k in class_keys if k)
        cleared = 0
        blocked = 0
        for ck in affected_classes:
            applicable_class = None if ck == 'common' else ck
            products_to_clear = Product.query.filter_by(city_id=city_id, applicable_class=applicable_class).all()
            for p in products_to_clear:
                from app.models import OrderItem, Order
                pending = OrderItem.query.join(Order).filter(
                    OrderItem.product_id == p.id, Order.status == 'pending'
                ).first()
                if pending:
                    blocked += 1
                    continue
                StockLog.query.filter_by(product_id=p.id).delete()
                db.session.delete(p)
                cleared += 1
        db.session.commit()
        if blocked:
            flash(f'Cleared {cleared} old item(s). {blocked} could not be cleared (part of a pending online order).')
        else:
            flash(f'Cleared {cleared} old item(s) from the affected classes.')

    created_count = 0
    for i in range(len(names)):
        name = names[i].strip() if i < len(names) else ''
        if not name:
            continue

        class_key = class_keys[i] if i < len(class_keys) else ''
        applicable_class = None if class_key == 'common' else (class_key or None)
        subject = subjects[i].strip() if i < len(subjects) else ''
        publisher = publishers[i].strip() if i < len(publishers) else ''
        is_elective = (elective_flags[i] == 'true') if i < len(elective_flags) else False

        gst_raw = gst_percents[i].strip() if i < len(gst_percents) else 'exempt'
        gst_percent, is_gst_exempt = parse_gst(gst_raw)

        try:
            price = float(prices[i]) if i < len(prices) and prices[i] else 0.0
        except ValueError:
            price = 0.0

        try:
            qty = int(quantities[i]) if i < len(quantities) and quantities[i] else 0
        except ValueError:
            qty = 0

        existing = Product.query.filter_by(city_id=city_id, name=name, applicable_class=applicable_class).first()
        if existing:
            existing.price = price
            existing.gst_percent = gst_percent
            existing.is_gst_exempt = is_gst_exempt
            existing.subject = subject or None
            existing.publisher_name = publisher or None
            existing.is_elective = is_elective
            product = existing
        else:
            product = Product(
                name=name, hsn_number='0000', gst_percent=gst_percent, is_gst_exempt=is_gst_exempt,
                price=price, product_type='textbook', applicable_class=applicable_class,
                publisher_name=publisher or None, subject=subject or None,
                city_id=city_id, stock_quantity=0, is_elective=is_elective
            )
            db.session.add(product)

        created_count += 1

    db.session.commit()
    flash(f'Imported/updated {created_count} item(s) into inventory')
    return redirect(url_for('inventory.class_list'))