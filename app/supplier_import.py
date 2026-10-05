import re
from datetime import datetime
from flask import Blueprint, render_template, redirect, url_for, request, flash
from flask_login import login_required, current_user
from openpyxl import load_workbook
from app import db
from app.models import Product, StockLog, Supplier, PurchaseRecord, PurchaseRecordItem
from app.inventory import store_staff_required, parse_gst
from app.city_utils import get_active_city_id
import pdfplumber
from app.activity_log import log_activity
from app.class_options import CLASS_OPTIONS
from app.models import Supplier, PurchaseRecord, PurchaseRecordItem, StockLog

supplier_import = Blueprint('supplier_import', __name__, url_prefix='/supplier-import')

HEADER_KEYWORDS = {
    'name': ['description', 'particular', 'title', 'goods', 'item name', 'product name'],
    'hsn': ['hsn', 'sac'],
    'qty': ['qty', 'quantity'],
    'price': ['rate', 'price'],
    'amount': ['taxable amt', 'amount', 'value'],
    'gst': ['gst rate', 'gst', 'tax rate'],
}
NOTEBOOK_KEYWORDS = ['notebook', 'register', 'copy', 'diary', 'line', 'square', 'spiral', 'file sheet', 'drawing']
SUBJECT_KEYWORDS = ['hindi', 'english', 'maths', 'math', 'science', 'evs', 'social', 'marathi', 'sanskrit', 'french', 'computer', 'theme']
ROMAN_MAP = {'i': 1, 'ii': 2, 'iii': 3, 'iv': 4, 'v': 5, 'vi': 6, 'vii': 7, 'viii': 8, 'ix': 9, 'x': 10, 'xi': 11, 'xii': 12}


def find_header_row(ws, max_scan=10):
    for row_idx in range(1, max_scan + 1):
        row_values = [str(c.value).strip().lower() if c.value else '' for c in ws[row_idx]]
        joined = ' '.join(row_values)
        if any(k in joined for k in ['hsn', 'sac']) and any(k in joined for k in ['qty', 'quantity']):
            return row_idx, row_values
    return None, None


def map_columns(header_values):
    col_map = {}
    for key, keywords in HEADER_KEYWORDS.items():
        for idx, val in enumerate(header_values):
            if any(kw in val for kw in keywords):
                col_map[key] = idx
                break
    return col_map


def guess_class(name):
    match = re.search(r'\b(class|grade|level)\s*[-:]?\s*([ivx]+|\d+)\b', name, re.IGNORECASE)
    if match:
        val = match.group(2).lower()
        if val.isdigit():
            return val
        return str(ROMAN_MAP.get(val, ''))
    return ''


def guess_subject(name):
    lower = name.lower()
    for kw in SUBJECT_KEYWORDS:
        if kw in lower:
            return kw.capitalize()
    return ''


def guess_type(name):
    lower = name.lower()
    if any(kw in lower for kw in NOTEBOOK_KEYWORDS):
        return 'notebook'
    if guess_class(name):
        return 'textbook'
    return 'extra'

def parse_gst_from_invoice(gst_val):
    """Returns gst_percent as a number. Handles '18%', 0.18 (fraction), 18 (already %), and '9+9' (CGST+SGST split)."""
    if gst_val is None or str(gst_val).strip() == '':
        return None

    text = str(gst_val).strip()

    if '+' in text:
        parts = re.findall(r'[\d.]+', text)
        try:
            return round(sum(float(p) for p in parts), 2)
        except ValueError:
            return None

    text = text.replace('%', '').strip()

    try:
        value = float(text)
    except ValueError:
        return None

    if value <= 1:
        value = value * 100

    return round(value, 2)


def clean_duplicated_name(name):
    match = re.match(r'^(.*?)\s*&\s*\1$', name.strip(), re.IGNORECASE)
    if match:
        return match.group(1).strip()
    return name


def extract_rows_from_sheet(ws, header_row_idx, col_map):
    rows = []
    total_scanned = 0
    for row in ws.iter_rows(min_row=header_row_idx + 1, values_only=True):
        if row is None:
            continue
        name_idx = col_map.get('name')
        if name_idx is None or name_idx >= len(row):
            continue
        name_val = row[name_idx]
        qty_idx = col_map.get('qty')
        has_qty = qty_idx is not None and qty_idx < len(row) and row[qty_idx] not in (None, '')
        if not has_qty:
            continue
        total_scanned += 1

        if not name_val or not str(name_val).strip():
            continue
        name = clean_duplicated_name(str(name_val).strip().replace('\n', ' '))
        if name.lower() in ('total', 'grand total', 'sub total'):
            continue

        def get_val(key):
            idx = col_map.get(key)
            if idx is None or idx >= len(row):
                return None
            return row[idx]

        qty_val = get_val('qty')
        hsn_val = get_val('hsn')
        price_val = get_val('price')
        amount_val = get_val('amount')
        gst_val = get_val('gst')

        try:
            qty_clean = re.sub(r'[^\d.]', '', str(qty_val)) if qty_val else ''
            qty = int(float(qty_clean)) if qty_clean else 0
        except (ValueError, TypeError):
            continue
        if qty <= 0:
            continue

        price = 0.0
        try:
            if price_val not in (None, ''):
                price = float(re.sub(r'[^\d.]', '', str(price_val)))
            elif amount_val not in (None, ''):
                amount = float(re.sub(r'[^\d.]', '', str(amount_val)))
                price = round(amount / qty, 2) if qty else 0.0
        except (ValueError, TypeError):
            price = 0.0

        hsn = str(hsn_val).strip() if hsn_val else ''

        gst_percent = parse_gst_from_invoice(gst_val)
        is_exempt = gst_percent is None or gst_percent == 0

        rows.append({
            'name': name, 'hsn': hsn, 'qty': qty, 'price': price,
            'gst_percent': gst_percent if gst_percent is not None else 0,
            'is_exempt': is_exempt,
            'section': guess_class(name), 'subject': guess_subject(name),
            'product_type': guess_type(name)
        })

    return rows, total_scanned


def parse_invoice_excel(file_stream):
    try:
        wb = load_workbook(file_stream, data_only=True)
    except Exception:
        return []

    all_rows = []
    for ws in wb.worksheets:
        header_row_idx, header_values = find_header_row(ws)
        if header_row_idx is None:
            continue

        col_map = map_columns(header_values)
        if 'name' not in col_map or 'qty' not in col_map:
            continue

        rows, total_scanned = extract_rows_from_sheet(ws, header_row_idx, col_map)

        # skip sheets where the "name" column is clearly broken (merged cell artifact) —
        # most rows have a qty but no matching name
        if total_scanned > 0 and len(rows) / total_scanned < 0.5:
            continue

        all_rows.extend(rows)

    return all_rows

def parse_invoice_pdf(file_stream):
    rows = []
    try:
        with pdfplumber.open(file_stream) as pdf:
            for page in pdf.pages:
                tables = page.extract_tables()
                for table in tables:
                    if not table or len(table) < 2:
                        continue

                    header_values = [str(c).strip().lower() if c else '' for c in table[0]]
                    col_map = map_columns(header_values)

                    if 'name' not in col_map or 'qty' not in col_map:
                        continue

                    for data_row in table[1:]:
                        if data_row is None:
                            continue
                        name_idx = col_map.get('name')
                        if name_idx is None or name_idx >= len(data_row):
                            continue
                        name_val = data_row[name_idx]
                        if not name_val or not str(name_val).strip():
                            continue
                        name = str(name_val).strip().replace('\n', ' ')
                        if name.lower() in ('total', 'grand total', 'sub total'):
                            continue

                        def get_val(key):
                            idx = col_map.get(key)
                            if idx is None or idx >= len(data_row):
                                return None
                            return data_row[idx]

                        qty_val = get_val('qty')
                        hsn_val = get_val('hsn')
                        price_val = get_val('price')

                        try:
                            qty_clean = re.sub(r'[^\d.]', '', str(qty_val)) if qty_val else ''
                            qty = int(float(qty_clean)) if qty_clean else 0
                        except (ValueError, TypeError):
                            continue
                        if qty <= 0:
                            continue

                        try:
                            price_clean = re.sub(r'[^\d.]', '', str(price_val)) if price_val else ''
                            price = float(price_clean) if price_clean else 0.0
                        except (ValueError, TypeError):
                            price = 0.0

                        hsn = str(hsn_val).strip() if hsn_val else ''

                        rows.append({
                            'name': name, 'hsn': hsn, 'qty': qty, 'price': price,
                            'section': guess_class(name), 'subject': guess_subject(name),
                            'product_type': guess_type(name)
                        })
    except Exception:
        return []

    return rows


@supplier_import.route('/upload', methods=['GET', 'POST'])
@login_required
@store_staff_required
def upload():
    if request.method == 'POST':
        file = request.files.get('invoice_file')
        if not file or file.filename == '':
            flash('Please choose a file')
            return redirect(url_for('supplier_import.upload'))

        filename_lower = file.filename.lower()

        if filename_lower.endswith(('.xlsx', '.xls')):
            parsed_rows = parse_invoice_excel(file)
        elif filename_lower.endswith('.pdf'):
            parsed_rows = parse_invoice_pdf(file)
        else:
            flash('File must be .xlsx, .xls, or .pdf')
            return redirect(url_for('supplier_import.upload'))

        if not parsed_rows:
            flash('Could not auto-read an item table from this file — you can still add items manually below')

        supplier_name = request.form.get('supplier_name', '').strip()
        invoice_number = request.form.get('invoice_number', '').strip()
        invoice_date = request.form.get('invoice_date', '').strip()

        return render_template('supplier_import/review.html', rows=parsed_rows,
                                supplier_name=supplier_name, invoice_number=invoice_number,
                                invoice_date=invoice_date, class_options=CLASS_OPTIONS)

    all_suppliers = Supplier.query.order_by(Supplier.name).all()
    return render_template('supplier_import/upload.html', suppliers=all_suppliers)


@supplier_import.route('/confirm', methods=['POST'])
@login_required
@store_staff_required
def confirm():
    city_id = get_active_city_id()

    supplier_name = request.form.get('supplier_name', '').strip()
    invoice_number = request.form.get('invoice_number', '').strip() or None
    invoice_date_str = request.form.get('invoice_date', '').strip()

    if not supplier_name:
        flash('Supplier name is required')
        return redirect(url_for('supplier_import.upload'))

    supplier = Supplier.query.filter_by(name=supplier_name).first()
    if not supplier:
        supplier = Supplier(name=supplier_name)
        db.session.add(supplier)
        db.session.flush()

    invoice_date = None
    if invoice_date_str:
        try:
            invoice_date = datetime.strptime(invoice_date_str, '%Y-%m-%d').date()
        except ValueError:
            invoice_date = None

    names = request.form.getlist('name')
    hsns = request.form.getlist('hsn')
    gst_percents = request.form.getlist('gst_percent')
    invoiced_qtys = request.form.getlist('invoiced_qty')
    received_qtys = request.form.getlist('received_qty')
    prices = request.form.getlist('price')
    sections = request.form.getlist('section')
    subjects = request.form.getlist('subject')
    product_types = request.form.getlist('product_type')
    publishers = request.form.getlist('publisher')

    purchase_record = PurchaseRecord(
        supplier_id=supplier.id, invoice_number=invoice_number,
        invoice_date=invoice_date, city_id=city_id, created_by_id=current_user.id
    )
    db.session.add(purchase_record)
    db.session.flush()

    created_count = 0
    for i in range(len(names)):
        name = names[i].strip()
        if not name:
            continue

        hsn = hsns[i].strip() if i < len(hsns) else ''
        gst_raw = gst_percents[i].strip() if i < len(gst_percents) else 'exempt'
        gst_percent, is_gst_exempt = parse_gst(gst_raw)
        invoiced_qty = int(invoiced_qtys[i]) if i < len(invoiced_qtys) and invoiced_qtys[i].isdigit() else 0
        received_qty = int(received_qtys[i]) if i < len(received_qtys) and received_qtys[i].isdigit() else invoiced_qty
        try:
            price = float(prices[i]) if i < len(prices) and prices[i] else 0.0
        except ValueError:
            price = 0.0
        section = sections[i].strip() if i < len(sections) else ''
        subject = subjects[i].strip() if i < len(subjects) else ''
        publisher = publishers[i].strip() if i < len(publishers) else ''
        product_type = product_types[i] if i < len(product_types) and product_types[i] in ('textbook', 'notebook', 'extra') else 'extra'
        applicable_class = section if section else None

        if received_qty <= 0:
            continue

        existing = Product.query.filter_by(city_id=city_id, name=name, hsn_number=hsn).first()

        if existing:
            existing.stock_quantity += received_qty
            if price > 0:
                existing.price = price
            product = existing
        else:
            product = Product(
                name=name, hsn_number=hsn or '0000', gst_percent=gst_percent, is_gst_exempt=is_gst_exempt,
                price=price, product_type=product_type, applicable_class=applicable_class,
                publisher_name=publisher or None, subject=subject or None,
                city_id=city_id, stock_quantity=received_qty, is_active=True
            )
            db.session.add(product)
            db.session.flush()

        db.session.add(StockLog(product_id=product.id, city_id=city_id, quantity_added=received_qty, created_by_id=current_user.id))
        db.session.add(PurchaseRecordItem(
            purchase_record_id=purchase_record.id, product_id=product.id, product_name=name,
            invoiced_qty=invoiced_qty, received_qty=received_qty, price=price,
            gst_percent=gst_percent, is_gst_exempt=is_gst_exempt
        ))
        created_count += 1

    db.session.commit()
    flash(f'Shipment confirmed — {created_count} item(s) added to inventory')
    return redirect(url_for('supplier_import.history'))


@supplier_import.route('/history')
@login_required
@store_staff_required
def history():
    city_id = get_active_city_id()
    records = PurchaseRecord.query.filter_by(city_id=city_id).order_by(PurchaseRecord.created_at.desc()).all()
    return render_template('supplier_import/history.html', records=records)


@supplier_import.route('/record/<int:record_id>')
@login_required
@store_staff_required
def view_record(record_id):
    record = PurchaseRecord.query.get_or_404(record_id)
    if record.city_id != get_active_city_id():
        flash('Access denied')
        return redirect(url_for('supplier_import.history'))
    return render_template('supplier_import/record_detail.html', record=record)

@supplier_import.route('/manual', methods=['GET', 'POST'])
@login_required
@store_staff_required
def manual_record():
    city_id = get_active_city_id()

    if request.method == 'POST':
        supplier_name = request.form.get('supplier_name', '').strip()
        invoice_number = request.form.get('invoice_number', '').strip() or None
        product_id = request.form.get('product_id')
        quantity = request.form.get('quantity', '').strip()

        if not supplier_name or not product_id or not quantity:
            flash('Supplier name, product and quantity are required')
            return redirect(url_for('supplier_import.manual_record'))

        try:
            quantity = int(quantity)
        except ValueError:
            flash('Quantity must be a number')
            return redirect(url_for('supplier_import.manual_record'))

        product = Product.query.get(int(product_id))
        if not product or product.city_id != city_id:
            flash('Invalid product selected')
            return redirect(url_for('supplier_import.manual_record'))

        supplier = Supplier.query.filter_by(name=supplier_name).first()
        if not supplier:
            supplier = Supplier(name=supplier_name)
            db.session.add(supplier)
            db.session.flush()

        purchase_record = PurchaseRecord(
            supplier_id=supplier.id, invoice_number=invoice_number,
            invoice_date=datetime.utcnow().date(), city_id=city_id, created_by_id=current_user.id
        )
        db.session.add(purchase_record)
        db.session.flush()

        db.session.add(PurchaseRecordItem(
            purchase_record_id=purchase_record.id, product_id=product.id, product_name=product.name,
            invoiced_qty=quantity, received_qty=quantity, price=product.price,
            gst_percent=product.gst_percent, is_gst_exempt=product.is_gst_exempt
        ))

        product.stock_quantity += quantity
        db.session.add(StockLog(
            product_id=product.id, city_id=city_id,
            quantity_added=quantity, created_by_id=current_user.id
        ))

        db.session.commit()

        log_activity('Recorded manual stock entry', f'{quantity} x "{product.name}" from "{supplier_name}" (Invoice: {invoice_number or "-"})')
        flash(f'Recorded: {quantity} x "{product.name}" from {supplier_name}')
        return redirect(url_for('supplier_import.manual_record'))

    from app.inventory import section_label, class_value_to_key
    products = Product.query.filter_by(city_id=city_id, is_active=True).order_by(Product.name).all()

    def sort_key(k):
        return (0, int(k)) if k.isdigit() else (1, k)

    products_sorted = sorted(products, key=lambda p: (
        sort_key(class_value_to_key(p.applicable_class) if class_value_to_key(p.applicable_class) not in ('notebook', 'extra') else 'stationary'),
        p.name
    ))

    for p in products_sorted:
        key = class_value_to_key(p.applicable_class)
        if key in ('notebook', 'extra'):
            key = 'stationary'
        p.display_label = section_label(key)

    suppliers = Supplier.query.order_by(Supplier.name).all()
    recent = PurchaseRecord.query.filter_by(city_id=city_id).order_by(PurchaseRecord.created_at.desc()).limit(10).all()
    return render_template('supplier_import/manual.html', products=products_sorted, suppliers=suppliers, recent=recent)