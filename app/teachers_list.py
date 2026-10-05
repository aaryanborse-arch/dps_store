from io import BytesIO
from flask import Blueprint, render_template, redirect, url_for, request, flash, send_file, jsonify
from flask_login import login_required, current_user
from functools import wraps
from openpyxl import Workbook, load_workbook
from app import db
from app.models import TeacherRecord, City
from app.inventory import store_staff_required
from app.city_utils import get_active_city_id

teachers_list_bp = Blueprint('teachers_list_bp', __name__, url_prefix='/teachers-list')


def super_admin_required(f):
    @wraps(f)
    def decorated(*args, **kwargs):
        if not current_user.is_authenticated or current_user.role != 'super_admin':
            flash('Access denied — Super Admin only')
            return redirect(url_for('auth.dashboard'))
        return f(*args, **kwargs)
    return decorated


@teachers_list_bp.route('/template')
@login_required
@super_admin_required
def download_template():
    wb = Workbook()
    ws = wb.active
    ws.append(['Teacher ID', 'Teacher Name'])
    ws.append(['T-101', 'Priya Sharma'])
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return send_file(buffer, as_attachment=True, download_name='teacher_list_template.xlsx',
                      mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


@teachers_list_bp.route('/', methods=['GET', 'POST'])
@login_required
@super_admin_required
def manage_teachers():
    all_cities = City.query.order_by(City.name).all()

    city_id_param = request.args.get('city_id') or request.form.get('city_id')
    city_id = int(city_id_param) if city_id_param else (all_cities[0].id if all_cities else None)

    if request.method == 'POST':
        if not city_id:
            flash('Select a school first')
            return redirect(url_for('teachers_list_bp.manage_teachers'))

        file = request.files.get('excel_file')
        if not file or file.filename == '':
            flash('Please choose an Excel file')
            return redirect(url_for('teachers_list_bp.manage_teachers', city_id=city_id))

        try:
            wb = load_workbook(file, data_only=True)
            ws = wb.active
        except Exception:
            flash('Could not read that file')
            return redirect(url_for('teachers_list_bp.manage_teachers', city_id=city_id))

        header_row = [str(c.value).strip().lower() if c.value else '' for c in ws[1]]

        def find_col(keywords):
            for idx, val in enumerate(header_row):
                if any(kw in val for kw in keywords):
                    return idx
            return None

        id_col = find_col(['teacher id', 'employee id', 'staff id', 'id'])
        name_col = find_col(['teacher name', 'name'])

        if id_col is None or name_col is None:
            flash(f'Could not find required columns. Found headers: {", ".join(h for h in header_row if h)}. Make sure columns are labeled "Teacher ID" and "Teacher Name".')
            return redirect(url_for('teachers_list_bp.manage_teachers', city_id=city_id))

        imported = 0
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or id_col >= len(row) or not row[id_col]:
                continue
            t_id = str(row[id_col]).strip()
            if t_id.endswith('.0'):
                t_id = t_id[:-2]
            t_name = str(row[name_col]).strip() if name_col < len(row) and row[name_col] else None
            if not t_id or not t_name:
                continue

            existing = TeacherRecord.query.filter_by(city_id=city_id, teacher_id=t_id).first()
            if existing:
                existing.teacher_name = t_name
            else:
                db.session.add(TeacherRecord(city_id=city_id, teacher_id=t_id, teacher_name=t_name))
            imported += 1

        db.session.commit()
        flash(f'Imported/updated {imported} teacher record(s)')
        return redirect(url_for('teachers_list_bp.manage_teachers', city_id=city_id))

    records = TeacherRecord.query.filter_by(city_id=city_id).order_by(TeacherRecord.teacher_name).all() if city_id else []
    selected_city = City.query.get(city_id) if city_id else None

    return render_template('teachers_list/manage.html', records=records, all_cities=all_cities, selected_city=selected_city)


@teachers_list_bp.route('/add', methods=['POST'])
@login_required
@super_admin_required
def add_teacher():
    city_id = request.form.get('city_id')
    teacher_id = request.form.get('teacher_id', '').strip()
    teacher_name = request.form.get('teacher_name', '').strip()

    if not city_id or not teacher_id or not teacher_name:
        flash('School, Teacher ID and Teacher Name are all required')
        return redirect(url_for('teachers_list_bp.manage_teachers'))

    city_id = int(city_id)

    existing = TeacherRecord.query.filter_by(city_id=city_id, teacher_id=teacher_id).first()
    if existing:
        existing.teacher_name = teacher_name
        flash(f'Updated existing record for Teacher ID {teacher_id}')
    else:
        db.session.add(TeacherRecord(city_id=city_id, teacher_id=teacher_id, teacher_name=teacher_name))
        flash(f'Added teacher "{teacher_name}" (Teacher ID {teacher_id})')

    db.session.commit()
    return redirect(url_for('teachers_list_bp.manage_teachers', city_id=city_id))


@teachers_list_bp.route('/lookup')
def lookup_teacher():
    city_id = request.args.get('city_id')
    teacher_id = request.args.get('teacher_id', '').strip()
    if not city_id or not teacher_id:
        return jsonify({'found': False})

    record = TeacherRecord.query.filter_by(city_id=int(city_id), teacher_id=teacher_id).first()
    if record:
        return jsonify({'found': True, 'name': record.teacher_name})
    return jsonify({'found': False})

@teachers_list_bp.route('/staff-add', methods=['GET', 'POST'])
@login_required
@store_staff_required
def staff_add_teacher():
    city_id = get_active_city_id()

    if request.method == 'POST':
        teacher_id = request.form.get('teacher_id', '').strip()
        teacher_name = request.form.get('teacher_name', '').strip()

        if not teacher_id or not teacher_name:
            flash('Teacher ID and Teacher Name are both required')
            return redirect(url_for('teachers_list_bp.staff_add_teacher'))

        existing = TeacherRecord.query.filter_by(city_id=city_id, teacher_id=teacher_id).first()
        if existing:
            existing.teacher_name = teacher_name
            flash(f'Updated existing record for Teacher ID {teacher_id}')
        else:
            db.session.add(TeacherRecord(city_id=city_id, teacher_id=teacher_id, teacher_name=teacher_name))
            flash(f'Added teacher "{teacher_name}" (Teacher ID {teacher_id})')

        db.session.commit()
        return redirect(url_for('teachers_list_bp.staff_add_teacher'))

    records = TeacherRecord.query.filter_by(city_id=city_id).order_by(TeacherRecord.teacher_name).all()
    return render_template('teachers_list/staff_add.html', records=records)