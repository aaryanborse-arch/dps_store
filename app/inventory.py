from flask import Blueprint, render_template, redirect, url_for, request, flash
from flask_login import login_required, current_user
from functools import wraps
from app import db
from app.models import Product, StockLog
from app.city_utils import get_active_city_id, get_active_city
from app.pdf_utils import to_roman
from app.class_options import CLASS_OPTIONS
from app.pdf_utils import to_roman
from app.activity_log import log_activity
from datetime import datetime
from flask import send_file
from app.models import StockLog


inventory = Blueprint('inventory', __name__, url_prefix='/inventory')


def store_staff_required(f):
    @wraps(f)
    def decorated(*args, **kwargs):
        if not current_user.is_authenticated or current_user.role != 'store_staff':
            flash('Access denied — Store Staff only')
            return redirect(url_for('auth.dashboard'))
        return f(*args, **kwargs)
    return decorated


def class_key_to_value(class_key):
    return None if class_key == 'common' else class_key


def class_value_to_key(applicable_class):
    return 'common' if applicable_class is None else applicable_class

def section_label(class_key):
    if class_key == 'common':
        return 'Common Items'
    elif class_key == 'notebook':
        return 'Notebook'
    elif class_key in ('extra', 'stationary'):
        return 'Stationary'
    else:
        from app.class_options import get_class_label
        return f'Class {get_class_label(class_key)}'


def parse_gst(gst_str):
    """Returns (gst_percent, is_exempt). Accepts 'exempt'/'na'/'nil' as exempt, else a number."""
    cleaned = gst_str.strip().lower()
    if cleaned in ('exempt', 'exempted', 'na', 'n/a', 'nil', 'none'):
        return 0.0, True
    return float(gst_str), False


@inventory.route('/')
@login_required
@store_staff_required
def class_list():
    city_id = get_active_city_id()
    all_products = Product.query.filter_by(city_id=city_id, is_active=True).order_by(Product.name).all()

    counts = {}
    for p in all_products:
        key = class_value_to_key(p.applicable_class)
        if key == 'extra':
            key = 'stationary'
        counts[key] = counts.get(key, 0) + 1

    def sort_key(k):
        return (0, int(k)) if k.isdigit() else (1, k)

    sorted_keys = sorted(counts.keys(), key=sort_key)
    sections = [{'key': k, 'label': section_label(k), 'count': counts[k]} for k in sorted_keys]

    return render_template('inventory/class_list.html', sections=sections, class_options=CLASS_OPTIONS)


@inventory.route('/new-class', methods=['POST'])
@login_required
@store_staff_required
def new_class():
    class_name = request.form.get('class_name', '').strip()
    if not class_name:
        flash('Enter a class name')
        return redirect(url_for('inventory.class_list'))
    return redirect(url_for('inventory.class_products', class_key=class_name))


@inventory.route('/class/<class_key>', methods=['GET', 'POST'])
@login_required
@store_staff_required
def class_products(class_key):
    city_id = get_active_city_id()
    applicable_class = class_key_to_value(class_key)
    label = section_label(class_key)

    if request.method == 'POST':
        name = request.form.get('name', '').strip()
        hsn = request.form.get('hsn_number', '').strip()
        gst_raw = request.form.get('gst_percent', '').strip()
        price = request.form.get('price', '').strip()
        product_type = request.form.get('product_type')
        stock = request.form.get('stock_quantity', '0').strip()
        publisher_name = request.form.get('publisher_name', '').strip() or None
        subject = request.form.get('subject', '').strip() or None

        if not all([name, hsn, gst_raw, price, product_type]):
            flash('Name, HSN, GST, Price and Type are required')
        else:
            try:
                gst_percent, is_exempt = parse_gst(gst_raw)
                initial_stock = int(stock) if stock else 0
                is_elective = request.form.get('is_elective') == 'on'
                if class_key == 'stationary':
                    save_class = 'stationary'
                elif class_key == 'notebook':
                    save_class = 'notebook'
                else:
                    save_class = applicable_class
                new_product = Product(
                    name=name,
                    year=request.form.get('year', '').strip() or None,
                    hsn_number=hsn,
                    group_name=request.form.get('group_name', '').strip() or None,
                    gst_percent=gst_percent,
                    is_gst_exempt=is_exempt,
                    price=float(price),
                    product_type=product_type,
                    applicable_class=save_class,
                    publisher_name=publisher_name,
                    subject=subject,
                    city_id=city_id,
                    stock_quantity=initial_stock,
                    is_elective=is_elective
                )
                db.session.add(new_product)
                db.session.flush()

                if initial_stock > 0:
                    db.session.add(StockLog(
                        product_id=new_product.id,
                        city_id=city_id,
                        quantity_added=initial_stock,
                        created_by_id=current_user.id
                    ))

                db.session.commit()
                flash(f'Product "{name}" added to {label}')
            except ValueError:
                flash('GST must be a number or "exempt". Price and Stock must be valid numbers')
        return redirect(url_for('inventory.class_products', class_key=class_key))

    if class_key == 'stationary':
        products = Product.query.filter(
            Product.city_id == city_id,
            Product.is_active == True,
            Product.applicable_class.in_(['extra', 'stationary'])
        ).order_by(Product.subject, Product.name).all()
    elif class_key == 'notebook':
        products = Product.query.filter_by(city_id=city_id, is_active=True, applicable_class='notebook').order_by(Product.subject, Product.name).all()
    elif applicable_class is None:
        products = Product.query.filter_by(city_id=city_id, is_active=True, applicable_class=None).order_by(Product.subject, Product.name).all()
    else:
        products = Product.query.filter_by(city_id=city_id, is_active=True, applicable_class=applicable_class).order_by(Product.subject, Product.name).all()

    existing_groups = sorted(set(p.group_name for p in products if p.group_name))

    display_rows = []
    seen_groups = set()
    for p in products:
        if p.group_name:
            if p.group_name in seen_groups:
                continue
            seen_groups.add(p.group_name)
            group_variants = [gp for gp in products if gp.group_name == p.group_name]

            for gp in group_variants:
                short = gp.name
                if short.startswith(gp.group_name):
                    short = short[len(gp.group_name):].lstrip(' -')
                gp.short_label = short or gp.name

            display_rows.append({'type': 'group', 'group_name': p.group_name, 'variants': group_variants})
        else:
            display_rows.append({'type': 'single', 'product': p})

    return render_template('inventory/class_products.html', products=products, display_rows=display_rows,
                            label=label, class_key=class_key, class_options=CLASS_OPTIONS, existing_groups=existing_groups)


@inventory.route('/<int:product_id>/edit', methods=['POST'])
@login_required
@store_staff_required
def edit_product(product_id):
    product = Product.query.get_or_404(product_id)
    city_id = get_active_city_id()

    if product.city_id != city_id:
        flash('Access denied — not your active school\'s product')
        return redirect(url_for('inventory.class_list'))

    old_price = product.price

    product.name = request.form.get('name', product.name).strip()
    product.year = request.form.get('year', '').strip() or None
    product.publisher_name = request.form.get('publisher_name', product.publisher_name or '').strip() or None
    product.subject = request.form.get('subject', product.subject or '').strip() or None
    product.hsn_number = request.form.get('hsn_number', product.hsn_number).strip()
    product.group_name = request.form.get('group_name', '').strip() or None

    gst_raw = request.form.get('gst_percent', '').strip()
    gst_percent, is_exempt = parse_gst(gst_raw)
    product.gst_percent = gst_percent
    product.is_gst_exempt = is_exempt

    new_price = float(request.form.get('price', product.price))
    product.price = new_price
    # stock_quantity is intentionally NOT editable here — only via Record Inward or Stock Correction
    product.is_elective = request.form.get('is_elective') == 'on'

    db.session.commit()

    changes = []
    if old_price != new_price:
        changes.append(f'Price: ₹{old_price} → ₹{new_price}')

    detail_text = f'"{product.name}"' + (' — ' + ', '.join(changes) if changes else ' — no price change')
    log_activity('Edited product', detail_text)

    flash(f'Product "{product.name}" updated')

    class_key = class_value_to_key(product.applicable_class)
    return redirect(url_for('inventory.class_products', class_key=class_key))


@inventory.route('/<int:product_id>/delete', methods=['POST'])
@login_required
@store_staff_required
def delete_product(product_id):
    from app.models import OrderItem, Order

    product = Product.query.get_or_404(product_id)
    city_id = get_active_city_id()

    if product.city_id != city_id:
        flash('Access denied — not your active school\'s product')
        return redirect(url_for('inventory.class_list'))

    class_key = class_value_to_key(product.applicable_class)

    pending = OrderItem.query.join(Order).filter(
        OrderItem.product_id == product.id, Order.status == 'pending'
    ).first()
    if pending:
        flash(f'Cannot delete "{product.name}" — it\'s part of a pending online order awaiting collection. Resolve that order first.')
        return redirect(url_for('inventory.class_products', class_key=class_key))

    product.is_active = False
    product.deactivated_at = datetime.utcnow()
    db.session.commit()

    log_activity('Deleted product (deactivated)', f'"{product.name}"')
    flash(f'Product "{product.name}" removed')
    return redirect(url_for('inventory.class_products', class_key=class_key))

@inventory.route('/class/<class_key>/delete', methods=['POST'])
@login_required
@store_staff_required
def delete_class_section(class_key):
    city_id = get_active_city_id()
    applicable_class = class_key_to_value(class_key)

    if applicable_class is None:
        products = Product.query.filter_by(city_id=city_id, applicable_class=None).all()
    else:
        products = Product.query.filter_by(city_id=city_id, applicable_class=applicable_class).all()

    for p in products:
        p.is_active = False

    db.session.commit()
    flash(f'Section deleted — {len(products)} product(s) removed from listing')
    return redirect(url_for('inventory.class_list'))

@inventory.route('/class/<class_key>/template')
@login_required
@store_staff_required
def download_class_template(class_key):
    from openpyxl import Workbook
    from io import BytesIO

    label = section_label(class_key)

    wb = Workbook()
    ws = wb.active
    ws.append(['Product', 'Class', 'Year', 'HSN Number', 'GST', 'Price', 'Existing Stock'])
    ws.append(['Example English Textbook', label, '2026-27', '4901', '18%', '250', '40'])
    ws.append(['Example Notebook (200 pg)', label, '2026-27', '', 'exempt', '60', '100'])

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    safe_label = label.replace(' ', '_')
    return send_file(buffer, as_attachment=True, download_name=f'{safe_label}_template.xlsx',
                      mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


@inventory.route('/class/<class_key>/import', methods=['POST'])
@login_required
@store_staff_required
def import_class_template(class_key):
    from openpyxl import load_workbook

    city_id = get_active_city_id()
    applicable_class = class_key_to_value(class_key)
    label = section_label(class_key)

    file = request.files.get('template_file')
    if not file or file.filename == '':
        flash('Please choose a filled template file')
        return redirect(url_for('inventory.class_products', class_key=class_key))

    try:
        wb = load_workbook(file, data_only=True)
        ws = wb.active
    except Exception:
        flash('Could not read that file')
        return redirect(url_for('inventory.class_products', class_key=class_key))

    header_row = [str(c.value).strip().lower() if c.value else '' for c in ws[1]]

    def find_col(keywords):
        for idx, val in enumerate(header_row):
            if any(kw in val for kw in keywords):
                return idx
        return None

    name_col = find_col(['product'])
    year_col = find_col(['year'])
    hsn_col = find_col(['hsn'])
    gst_col = find_col(['gst'])
    price_col = find_col(['price'])
    stock_col = find_col(['stock'])

    if name_col is None or price_col is None:
        flash('Could not find required columns (Product, Price) in the file')
        return redirect(url_for('inventory.class_products', class_key=class_key))

    created_count = 0
    updated_count = 0
    stock_ignored_count = 0

    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or name_col >= len(row) or not row[name_col]:
            continue

        name = str(row[name_col]).strip()
        if not name:
            continue

        year_val = str(row[year_col]).strip() if year_col is not None and year_col < len(row) and row[year_col] else None
        hsn = str(row[hsn_col]).strip() if hsn_col is not None and hsn_col < len(row) and row[hsn_col] else '0000'
        gst_raw = str(row[gst_col]).strip() if gst_col is not None and gst_col < len(row) and row[gst_col] else 'exempt'
        try:
            gst_percent, is_exempt = parse_gst(gst_raw)
        except (ValueError, TypeError):
            gst_percent, is_exempt = 0.0, True

        try:
            price = float(row[price_col]) if price_col < len(row) and row[price_col] not in (None, '') else 0.0
        except (ValueError, TypeError):
            price = 0.0

        try:
            stock = int(row[stock_col]) if stock_col is not None and stock_col < len(row) and row[stock_col] not in (None, '') else 0
        except (ValueError, TypeError):
            stock = 0

        save_class = 'stationary' if class_key == 'stationary' else ('notebook' if class_key == 'notebook' else applicable_class)

        existing = Product.query.filter_by(city_id=city_id, name=name, applicable_class=save_class).first()
        if existing:
            existing.hsn_number = hsn
            existing.gst_percent = gst_percent
            existing.is_gst_exempt = is_exempt
            existing.price = price

            old_stock = existing.stock_quantity
            if stock != old_stock:
                diff = stock - old_stock
                existing.stock_quantity = stock
                db.session.add(StockLog(
                    product_id=existing.id, city_id=city_id,
                    quantity_added=diff, created_by_id=current_user.id
                ))
                stock_ignored_count += 1  # reused as "stock corrected" counter below

            updated_count += 1
        else:
            new_product = Product(
                name=name, hsn_number=hsn, gst_percent=gst_percent, is_gst_exempt=is_exempt,
                price=price, product_type='textbook' if class_key not in ('stationary', 'notebook', 'common') else 'extra',
                applicable_class=save_class, city_id=city_id, stock_quantity=stock
            )
            db.session.add(new_product)
            created_count += 1

    db.session.commit()

    log_activity('Imported class template', f'{label}: {created_count} created, {updated_count} updated')

    msg = f'Imported into {label}: {created_count} new item(s) added, {updated_count} existing item(s) updated'
    if stock_ignored_count:
        msg += f'. {stock_ignored_count} item(s) had their stock corrected to match the sheet.'
    flash(msg)

    return redirect(url_for('inventory.class_products', class_key=class_key))